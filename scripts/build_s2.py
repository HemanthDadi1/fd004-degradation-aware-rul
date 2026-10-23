"""Build Supplementary Data S2 for Part II from the run folders. python build_s2.py ROOT OUT.xlsx"""
import re, sys
from pathlib import Path
import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT, OUT = Path(sys.argv[1]), sys.argv[2]
PROTO = {"IID": "IID", "DS_A": "DS-A", "DS_B": "DS-B"}


def describe(name):
    m = re.match(r"FD004_(?:(FD002)_)?(IID|DS_A|DS_B)_(?:W(\d+)|Rerun)(?:_(bilstm_attn|lstm_attn|bilstm_noattn|lstm_noattn))?_Results$", name)
    if not m or name == "FD004_DS_B_Results":
        return None
    ds, proto, w, arch = m.groups()
    return {"dataset": "FD002" if ds else "FD004", "protocol": PROTO[proto], "window": int(w) if w else 30,
            "architecture": arch or "bilstm_attn", "source": "Part I archive" if (w is None or (ds and w == "30")) else "Part II"}


rows, splits = [], []
for d in sorted(ROOT.glob("FD004_*_Results")):
    info = describe(d.name)
    if info is None or not (d / "convergence_check.csv").exists():
        continue
    restored = d / "convergence_check_restored_epoch.csv"
    if restored.exists():  # Keras ignores epochs 1-30 for early stopping, so judge convergence on the restored epoch
        c = pd.read_csv(restored, comment="#")
        c = c.rename(columns={"restored_epoch": "restored_epoch", "restored_ratio": "loss_ratio"})
        c = c[["seed", "restored_epoch", "loss_ratio", "converged"]] if "loss_ratio" in c else None
    else:
        c = None
    if c is None:
        raw = pd.read_csv(d / "convergence_check.csv")
        c = raw.rename(columns={"best_epoch": "restored_epoch", "ratio": "loss_ratio"})[["seed", "restored_epoch", "loss_ratio", "converged"]]
    c = c.copy()
    for k, v in info.items():
        c[k] = v
    c["run_folder"] = d.name
    rows.append(c)
    if info["source"] == "Part II" and info["dataset"] == "FD002" and (d / "engine_split.csv").exists():
        s = pd.read_csv(d / "engine_split.csv"); s.insert(0, "protocol", info["protocol"]); splits.append(s)
conv = pd.concat(rows)
conv["collapsed (ratio > 0.95)"] = conv.loss_ratio > 0.95
conv = conv[["dataset", "protocol", "window", "architecture", "source", "seed", "restored_epoch", "loss_ratio", "converged",
             "collapsed (ratio > 0.95)", "run_folder"]].sort_values(["dataset", "architecture", "protocol", "window", "seed"])

p2 = conv[conv.source == "Part II"]
summary = (p2.groupby(["dataset", "architecture", "protocol", "window"])
           .agg(runs=("seed", "size"), converged=("converged", "sum"), collapsed=("collapsed (ratio > 0.95)", "sum"),
                max_ratio=("loss_ratio", "max")).reset_index())
flagged = conv[(~conv.converged.astype(bool)) | (conv.loss_ratio > 0.6)].sort_values("loss_ratio", ascending=False)

e1 = pd.read_csv(ROOT / "FD004_P2E1_vs_P2E8_Results/convergence_check_restored_epoch.csv", comment="#")
sheets = {
    "Convergence summary": summary,
    "Convergence all runs": conv,
    "Flagged runs": flagged,
    "P2E1 vs P2E8 convergence": e1,
    "FD002 splits": pd.concat(splits) if splits else pd.DataFrame(),
    "Matched 30v60 FD004 by seed": pd.read_csv(ROOT / "FD004_Paper2_Matched_30v60/matched_metrics_by_seed.csv"),
    "Matched 30v60 FD002 by seed": pd.read_csv(ROOT / "FD004_FD002_Matched_30v60/matched_metrics_by_seed.csv"),
    "Ensembles": pd.read_csv(ROOT / "FD004_Paper2_Ensembles/ensemble_metrics.csv"),
    "Ensemble bootstrap": pd.read_csv(ROOT / "FD004_Paper2_Ensembles/ensemble_bootstrap_60_minus_30.csv"),
    "Equal width": pd.read_csv(ROOT / "FD004_Paper2_EqualWidth/equal_width_30v60.csv"),
    "Factorial by seed": pd.read_csv(ROOT / "FD004_Paper2_Factorial_Sweep/factorial_matched_by_seed.csv"),
    "Sweep by seed": pd.read_csv(ROOT / "FD004_Paper2_Factorial_Sweep/sweep_matched_by_seed.csv"),
}
readme = [
    ("Supplementary Data S2", "Part II: Does a Longer Temporal Context Help? Companion to Supplementary Data S1 of Part I."),
    ("Engine splits (FD004)", "The FD004 IID, DS-A and DS-B splits and the DS-B envelope thresholds are those of Supplementary Data S1 (Part I); every Part II run uses them."),
    ("FD002 splits", "Sheet 'FD002 splits': engine ids by split for the three FD002 protocols at 60 cycles (the 30-cycle FD002 runs of Part I use the same engines)."),
    ("Convergence", "A run counts as converged if the restored model's internal-validation loss is at most 0.60 times the loss of constant training quantiles. Keras ignores epochs 1-30 for early stopping, so 'restored_epoch' is the best epoch from 31 on. A run with loss_ratio above 0.95 predicts the constant quantiles ('collapsed')."),
    ("Flagged runs", "Every run that missed the 0.60 threshold or has a ratio above 0.60: the near miss (IID 60-cycle seed 123) and the collapsed runs. All are reported in the paper; Tables 11-12 leave collapsed seeds out of the means."),
    ("Matched sheets", "Per-seed metrics on identical endpoints (same seed, engine, endpoint cycle). 'Factorial by seed' is scored on the endpoints a 60-cycle window can score, 'Sweep by seed' on those a 90-cycle window can score."),
    ("Ensembles, Equal width", "Five-seed ensembles (quantiles averaged, then calibrated) and Window CQR level curves read at equal mean width (draft Tables 8-9, Figure 3)."),
    ("Source files", "Run folders FD004_*_Results with models, histories, row-level predictions and SHA-256 manifests: https://github.com/HemanthDadi1/fd004-degradation-aware-rul [DOI to be added]."),
]
with pd.ExcelWriter(OUT, engine="openpyxl") as w:
    pd.DataFrame(readme, columns=["item", "description"]).to_excel(w, "README", index=False)
    for name, df in sheets.items():
        df.to_excel(w, name[:31], index=False)
    for ws in w.book.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True)
        for i, col in enumerate(ws.columns, 1):
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, max(len(str(c.value or "")) for c in list(col)[:60]) + 2))
        ws.freeze_panes = "A2"
print(summary.to_string(index=False)); print(flagged[["dataset", "architecture", "protocol", "window", "seed", "loss_ratio", "converged"]].to_string(index=False))
print(len(conv), "runs;", len(p2), "from Part II; sheets:", len(sheets) + 1)
